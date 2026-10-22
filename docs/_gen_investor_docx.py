# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def set_run_font(run, size=11, bold=False, color=None, name="微软雅黑"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def add_heading_cn(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            name="微软雅黑",
        )
    return p


def add_para(text, size=11, bold=False, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_bullets(items, size=11):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3


t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(t.add_run("Diary One"), size=26, bold=True)

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(st.add_run("长期个人记忆系统 · 投资人沟通材料"), size=16, bold=True)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(
    meta.add_run(
        "NEXLEAP 跃迁光年  |  V0.1  |  2026年10月\n"
        "综合：《AI摄像耳机商业分析报告》《Diary One App 产品定义书 PRD V0.1》\n"
        "及《长期个人记忆系统建设策略——研究需求定义 V0.1》"
    ),
    size=9,
    color=RGBColor(0x55, 0x55, 0x55),
)

add_para(
    "说明：本文面向投资人沟通，阐述品类判断、产品形态、记忆库护城河、AI 能力、"
    "用户数据与世界模型关系，以及 Chapter / Diary 的产品落点。"
    "研究结论与商业假设分离标注；不以未验证指标冒充已完成验收。",
    size=9,
)

add_heading_cn("一、一句话与投资命题", 1)
add_para(
    "Diary One 帮助重视亲身经历、却不愿持续拍摄和整理的人，把经过授权的稀疏生活素材，"
    "转成真实、简洁、值得回看并能长期找回的个人记忆。"
)
add_para(
    "投资命题：品类正在从「纯听记忆设备」与「智能眼镜心智」之外，开辟第一人称「看 + 听」"
    "的个人记忆入口；耳机形态消解眼镜的佩戴与视力约束。真正可规模化的壁垒，不在单次生成的"
    "漂亮日记，而在每位用户可持续积累、更新与使用的长期个人记忆库，以及由此形成的授权数据飞轮。"
    "硬件是数据采集入口，App 的 Chapter / Diary 是价值交付面，记忆库与 AI 理解能力是核心资产。"
)

add_heading_cn("二、市场与品类判断（商业分析要点）", 1)
add_para(
    "智能眼镜已验证「拍摄 + AI」需求（全球 AI 眼镜出货高速增长），但面临视力、镜片、脸型、"
    "重量、外观等结构性约束。耳机作为每日既有佩戴设备，具备声音交互与形态渗透优势；产业侧 "
    "Meta、苹果、三星、字节等均在布局摄像耳机，但截至报告窗口多数尚未规模上市，存在创业公司落地窗口。"
)
add_para(
    "纯听觉记忆设备路径已被证伪或收编（如 Limitless、Bee 等事件所提示）：行业共识转向"
    "「看 + 听」的记忆系统。挂坠形态 Looki 等证明了约 $199–250 价位与资本兴趣，也暴露隐私社交阻力。"
    "因此产品必须坚持：隐私指示与授权优先于功能堆砌；架构上「采集在端侧、理解在云端」，"
    "以应对耳机电池与算力硬约束。"
)
add_bullets(
    [
        "形态选择：AI 摄像耳机 = AI 记忆硬件入口，服务长期个人记忆，而非单纯拍摄玩具。",
        "价格带参考：带摄像头个人 AI 硬件众筹与竞品多落在约 $199–299。",
        "最大风险：隐私立法与社会接受度；续航与温升决定采集策略必须稀疏、可授权、可关闭。",
    ]
)

add_heading_cn("三、产品形态：从素材到 Chapter / Diary", 1)
add_para(
    "产品信息架构以 Diary、Memory、Device、Account 四页为主。当前阶段（PRD M2）成功定义是："
    "用户主动选入照片/视频 → 手机 Evidence 编排 → 自有云算法与模型 → 结果回写 Diary；"
    "硬件未接入时诚实显示未连接，不伪造设备状态。"
)

add_heading_cn("3.1 处理链条", 2)
add_para(
    "原始素材 → 证据与观察事实 → 经历片段 → Chapter / Aha → Diary。"
    "中间结果与最终叙事分层：观测与叙事分离；宁少勿编；失败可恢复；未获云同意不上传。"
)

add_heading_cn("3.2 最终交付物：Chapter 与 Diary", 2)
add_bullets(
    [
        "Chapter（记忆卡）：当日或片段级的可回看单元，绑定证据与可追溯来源，是用户「认出自己经历」的主载体。",
        "Aha（小记忆）：轻量、可检索的补充记忆点，服务细节找回而非长文堆砌。",
        "Diary：按日组织的回顾界面，含今日状态、日历选日、历史分页；是长期使用习惯的主入口。",
        "Memory：字面搜索、Captured by You 原片、Topics 与 Preferences；支撑找回与偏好表达。",
        "Share：日级 Daily Recap、拼贴与插画等，在用户预览与授权后外发。",
    ]
)
add_para(
    "产品成功标准不是多生成文件或堆字数，而是：低负担记录、可信回忆、有用的长期积累、用户掌控。"
    "导入素材 ≠ 任意上传或训练授权。"
)

add_heading_cn("四、核心壁垒：长期个人记忆库", 1)
add_para(
    "团队判断：记忆库是 Diary One 的核心壁垒；公司定位必然走向「AI 记忆硬件 + 个人记忆系统」。"
    "单次 Chapter 可被竞品模仿，但「按用户隔离、可纠错、可找回、随年月变厚」的记忆资产与理解模型，"
    "构成数据与体验护城河。需同步启动研究，覆盖 AI 硬件竞品、GitHub 开源与学术论文。"
)

add_heading_cn("4.1 任务目标（研究需求定义）", 2)
add_para(
    "在尽量减少用户记录与整理负担的前提下，为大量用户分别建立可持续积累、更新和使用的长期个人记忆库，"
    "支撑三类核心体验："
)
add_bullets(
    [
        "自动记忆：从已获信息中保留对该用户有长期价值的内容。",
        "自然语言找回：用日常表达找回经历、细节、素材或一段时间的变化。",
        "个性化呈现：在合适场景呈现用户更想记住、更愿回看的内容。",
    ]
)

add_heading_cn("4.2 必须区分的记忆内容", 2)
add_bullets(
    [
        "经历与事实：记过什么、何时、原始依据是什么。",
        "跨时间关联：是否同一活动/对象/地点；如何避免错误关联。",
        "行为习惯：反复出现是否足以称为习惯。",
        "兴趣与偏好：明确表达 vs 行为推断。",
        "个人意义与记忆意愿：为何留下；一次性事件是否仍值得长期保存。",
    ]
)
add_para(
    "必须区分「观察到的事实」「用户明确表达」「系统推断」。"
    "系统不能因自身生成内容被反复引用，就将其当作新增事实依据。"
)

add_heading_cn("4.3 核心能力要求（摘要）", 2)
add_bullets(
    [
        "生命周期：提取、关联、去重、保存、更新、纠错、过期、遗忘与删除；重要记忆可说明来源、时间、证据状态与更新关系。",
        "习惯与偏好：处理冷启动、稀疏、中断与偏好变化；「没记到」≠「没发生」；明确反馈优先于行为推断冲突时的规则，并允许用户纠正或关闭。",
        "自然语言搜索：模糊找回、细节查找、跨时间汇总、变化比较、习惯询问；证据不足时明确说明。",
        "无感与个性化：值得记住什么 / 此刻适合呈现什么 / 怎样呈现；避免只重复已点击内容；无价值时可选择不生成、不打扰。",
    ]
)

add_heading_cn("4.4 规模化要求", 2)
add_bullets(
    [
        "用户隔离：存储、检索、缓存与生成按权限隔离。",
        "差异适应：不同生活方式、语言、记录频率。",
        "长期运行与成本性能：随历史增长仍可控；区分用户总量、DAU、并发查询、单用户日数据量与保存年限。",
        "可靠性与用户控制：重复上传、中断补传、版本升级；查看、纠错、删除、导出与个性化开关。",
    ]
)

add_heading_cn("五、AI 能力与端云分工", 1)
add_para(
    "受耳机电池与算力约束，可行架构是：端侧负责授权采集、压缩、去重、最小 Evidence 与本地账本；"
    "云端负责理解、编排、生成 Chapter / Aha 及后续记忆库维护。"
    "App 不内嵌模型密钥；服务端统一调用大模型（如通义等）并执行 fail-closed 准入。"
)
add_bullets(
    [
        "端侧：相册/设备素材导入、SHA256 去重、视频抽帧与质量门控、语义路由、Evidence 规范化、SQLite 账本、OTP 会话鉴权。",
        "云端 BFF / Worker：幂等 day-revisions、任务轮询、对象存储、代理算法编排（X02 等）、分享渲染等。",
        "算法原则：观测与叙事分离；未授权不上云；结果可追溯到证据；失败可重试，不以 mock 冒充生产。",
        "演进：从「当日精选日记」扩展到跨日关联、习惯与偏好、自然语言找回与个性化呈现；研究阶段不预先锁死数据库与模型栈。",
    ]
)

add_heading_cn("六、用户数据对世界模型的贡献——投资人结论", 1)
add_para(
    "结论（团队视角）：Diary One 用户数据对「世界模型」的贡献，首先不是无差别喂给通用大模型的公网语料，"
    "而是一类稀缺的、带时间与证据约束的第一人称生活轨迹数据；其投资价值在于「授权个人记忆飞轮」与"
    "「可验证的具身/日常场景分布」，而非「把用户日记直接拿去训练通用世界模型」的口号。"
)

add_heading_cn("6.1 为什么有独特价值", 2)
add_bullets(
    [
        "视角稀缺：第一人称、连续时间、跨日关联的真实经历分布，与互联网图文/影视第三人称数据互补。",
        "证据约束：素材—证据—叙事分层，适合研究「可追溯、可纠错」的记忆与推理，而非黑箱编造。",
        "长期动态：习惯与偏好会变化，逼近真实世界中的非平稳分布，对持续学习与遗忘机制有研究意义。",
        "硬件入口：AI 记忆硬件扩大采集覆盖后，数据从「偶发相册」走向「稀疏但真实的日常」，规模效应取决于授权与留存，而非一次性爬取。",
    ]
)

add_heading_cn("6.2 明确边界（对投资人诚实）", 2)
add_bullets(
    [
        "用户日记与原片默认属于用户；导入 ≠ 训练授权。任何用于模型改进的用途必须显式同意、可退出、可删除。",
        "未记录到的生活不能假装被模型「理解」；记忆库不声称全知世界模型。",
        "短期商业护城河优先落在：更高找回成功率、更可信 Chapter、更低整理负担、更强多用户隔离与合规——这些可产品化、可度量。",
        "对通用世界模型的贡献是中长期期权：在合规前提下，授权脱敏的结构化经历分布、失败与纠错信号、跨时间关联样本，可反哺感知—记忆—检索类模型；不应夸大为已具备可交易的「世界模型数据集」资产。",
    ]
)

add_heading_cn("6.3 一句话给投资人", 2)
add_para(
    "我们卖的是「AI 记忆硬件 + 可信个人记忆系统」；用户数据在严格授权下，既能加深产品护城河（个人记忆库），"
    "又为面向真实日常的具身/世界模型研究提供稀缺的第一人称、带证据的时间序列分布——"
    "后者是战略期权，前者是当下必须打穿的壁垒。"
)

add_heading_cn("七、研究启动与交付物", 1)
add_para(
    "记忆库建设以研究需求定义 V0.1 为纲，同步从 AI 硬件竞品、开源项目与论文中吸收可复用机制。"
    "研究交付不替代开发上线，但须支撑产品取舍："
)
add_bullets(
    [
        "现有链条信息地图（来源、可用性、证据质量、保留价值）。",
        "长期记忆内容与生命周期定义。",
        "关键体验定义：自动记忆、自然语言找回、习惯查看、个性化回顾与纠错。",
        "建设策略比较与分阶段实施范围。",
        "验证与验收：区分技术可运行、结果可信、用户价值成立；重视延迟回看，而非条数与点击。",
    ]
)

add_heading_cn("八、阶段路径（与 PRD 对齐）", 1)
add_bullets(
    [
        "M2 小规模真链路：主动选入 → 手机 → 自有云 → Diary Chapter（当前工程主线）。",
        "M3 完整非硬件内测：原片、分享、偏好、身份与恢复等补齐。",
        "M4 设备闭环：真实采集、配对、同步与稳定协议。",
        "M5 发布候选：正式账号、数据权利、成本运营与商店材料。",
        "并行：长期记忆库研究与试点能力（跨日关联、自然语言找回、个性化）按研究验收逐步嵌入 Memory / Diary。",
    ]
)

add_heading_cn("九、总结", 1)
add_para(
    "品类窗口在「看 + 听」的个人记忆硬件；Diary One 以可信 Chapter / Diary 交付日常价值，"
    "以长期个人记忆库构建壁垒，以端云 AI 在授权范围内完成理解与呈现。"
    "用户数据对世界模型的意义，是合规前提下的第一人称证据化生活分布与记忆飞轮，"
    "而不是无授权的训练权主张。下一步：打穿 M2 真链路与多人云底座，同步启动记忆库研究，"
    "用可验证的找回与回看价值对话用户与投资人。"
)

foot = doc.add_paragraph()
set_run_font(
    foot.add_run(
        "— 文档结束 —\n"
        "依据文件：AI摄像耳机商业分析报告；Diary_One_App_产品定义书_PRD_V0.1；"
        "长期个人记忆系统建设策略研究需求定义 V0.1；团队内部讨论（记忆库壁垒与世界模型贡献）。"
    ),
    size=8,
    color=RGBColor(0x66, 0x66, 0x66),
)

out = Path(r"E:\AI耳机\diary-one\docs\Diary_One_投资人沟通材料_记忆库与护城河_V0.1.docx")
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(out))
# also copy to workspace root for easy find
out2 = Path(r"E:\AI耳机\Diary_One_投资人沟通材料_记忆库与护城河_V0.1.docx")
doc.save(str(out2))
print(out)
print(out2)
print("paragraphs", len(doc.paragraphs))
print("chars", sum(len(p.text) for p in doc.paragraphs))
