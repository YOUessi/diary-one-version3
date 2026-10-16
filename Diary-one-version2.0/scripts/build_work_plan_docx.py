# -*- coding: utf-8 -*-
"""Generate Diary One algorithm + software work plan Word for multi-user beta."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "docs" / "Diary_One_算法与软件工作计划_多人内测.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)
    return p


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=9)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Diary One 算法与软件工作计划")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("面向多人内测（约几十人）· 对照当前代码与手册")
    set_run_font(r, size=12, color=RGBColor(0x55, 0x55, 0x55))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("版本日期：2026-10-08  |  仓库：Nexleap-Tech/Diary_one  |  配套：INTERNAL_BETA_FULL_GUIDE / EXPO_GO_FULL_TEST_GUIDE")
    set_run_font(r, size=9, color=RGBColor(0x77, 0x77, 0x77))

    add_para(
        doc,
        "本文用汇报口径说明：算法侧与软件侧各自要做什么、卡在哪、多人内测风险在哪、按什么顺序关项。"
        "目标不是「功能清单好看」，而是「几十个内测用户各自稳定拿到自己的 chapter」。",
    )

    # ---- 1 ----
    add_heading(doc, "一、目标与边界", 1)
    add_heading(doc, "1.1 内测成功长什么样", 2)
    add_bullets(
        doc,
        [
            "内测员用自己的邮箱登录（OTP），上传本人真实照片/短视频。",
            "系统完成：本地轻处理 → 上云算法筛选/出稿 → 在 Diary 看到本人当天 chapter。",
            "不同用户互不串数据；失败有可读错误；百炼账单可控。",
            "分发方式：TestFlight（优先）或固定 Expo 入口；有人值守失败任务。",
        ],
    )
    add_heading(doc, "1.2 本轮明确不做（不挡内测）", 2)
    add_bullets(
        doc,
        [
            "真耳机 BLE / 真电量 / 真自动采集硬件链。",
            "全员本机 Apple Vision（Expo Go 无原生 Vision；Dev Client/EAS 另轨）。",
            "本地自训小模型正式上云替换百炼（adapter 可后换，不挡开测）。",
            "精细订阅计费、OSS 替换 MinIO、完整监控大盘（包后加深）。",
        ],
    )
    add_heading(doc, "1.3 当前真实水位（别自欺）", 2)
    add_table(
        doc,
        ["能力", "状态", "对多人内测的含义"],
        [
            ["Expo Go 四 Tab + 入门/问卷", "可测", "UI/产品路径可验收"],
            ["相册→Bundle→BFF→chapter", "本机可测", "同 Wi‑Fi 几台手机够用，不够几十人"],
            ["Docker 量产栈代码", "已齐", "需部署到公网才服务外人"],
            ["真千问 / 关 mock", "看 Key 与配置", "不关 mock 则质量不可信"],
            ["MobileCLIP 本机权重", "可部署", "仅 BFF 所在机；照片主路径可不依赖"],
            ["公网 HTTPS BFF", "未立", "多人内测主卡点"],
            ["每用户日配额", "待补丁", "防打爆百炼"],
            ["TestFlight / EAS preview", "配置在，未发", "普通人分发依赖它"],
        ],
    )

    # ---- 2 ----
    add_heading(doc, "二、整体链路（算法 × 软件怎么分工）", 1)
    add_para(doc, "一条用户故事拆开看：", bold=True)
    add_bullets(
        doc,
        [
            "手机（软件）：选素材 → 校验 → SHA 去重 →（视频）探针抽帧 → 调 BFF 拿向量或降级 → 规则选 1/3/5/10 帧 → 压成 768×1024 JPEG → 打 V04 Bundle → 上传。",
            "BFF/Worker（软件基建）：鉴权、落盘、排队、调 X02、回写 job 状态；可选 MobileCLIP 出 512 维向量。",
            "X02 + 百炼（算法）：对 Bundle 做观察/筛选/出稿（V1.75 合同），产出 chapters/ahas。",
            "手机（软件）：轮询 job → map 到 Diary → 用户看到 chapter；分享走拼贴/插画/Recap。",
        ],
    )
    add_para(
        doc,
        "关键纪律：本地轻处理（去重/压图/选帧）≠ 云端大量素材筛选出稿；Sync ≠ Cloud Processing；Demo 小图经常空日，质量验收必须用真实日常图。",
    )

    # ---- 3 ----
    add_heading(doc, "三、算法工作计划", 1)
    add_heading(doc, "3.1 已具备 / 可复用", 2)
    add_bullets(
        doc,
        [
            "V1.75 筛选与出稿合同：与 eval-web / App Bundle（V04 letterbox 768×1024）对齐。",
            "云端通道脚本：open-cloud-bulk-channel.ps1 / smoke-cloud-pack.ps1（-Quality -RequireGrounded）。",
            "百炼 DashScope 作为订阅云模型；本地自训模型为可切换 adapter（另轨）。",
            "视频侧：MobileCLIP-S0 图像向量（512-d）已可在 Windows BFF 加载；场景标签包可选。",
        ],
    )
    add_heading(doc, "3.2 开多人内测前必须关的算法项", 2)
    add_table(
        doc,
        ["序号", "事项", "细节", "验收"],
        [
            [
                "A1",
                "真 qwen 常驻",
                "DASHSCOPE_API_KEY 写入运行环境；X02 :8787 与 BFF 同生命周期",
                "health/x02=up；任务非 fixture",
            ],
            [
                "A2",
                "关闭 mock 质量门",
                "ALLOW_MOCK_PREVIEW=false；禁止空日 mock_fallback 当成功",
                "真图 pack 有 grounded chapters",
            ],
            [
                "A3",
                "至少 1 份真实日 pack",
                "PC -RequireGrounded 或手机 Lab B 多张日常图",
                "报告无 dev_lite_fallback、非空 chapters",
            ],
            [
                "A4",
                "空日与失败语义统一",
                "证据不足应诚实失败/空日，不装有记忆",
                "文案与 Active tasks 一致、可解释",
            ],
            [
                "A5",
                "视频策略定稿",
                "公网机装 MobileCLIP 或明确「照片主验、视频可降级」",
                "内测说明写清；Lab C 可核对 analyzer",
            ],
        ],
    )
    add_heading(doc, "3.3 算法细节问题（容易踩坑）", 2)
    add_bullets(
        doc,
        [
            "Demo JPEG / 纯色图：常触发 NO_GROUNDED_MEMORY / EMPTY_DAY——这是证据问题，不是「算法没接上」。",
            "本地 eval 通过 ≠ 云端通过：云端多一层上传、MinIO、Worker、幂等、mock 开关。",
            "感知批上限：约 3 assets / 18 frames / 4 MiB；单日约 240 assets——大批量要切日或分 job。",
            "MobileCLIP 无权重或未重启 BFF：App 回落直方图 degraded，只能证明规则链路，不能当语义质量验收。",
            "场景标签（ml-mobileclip）装不上时：仍可用纯图像向量做选帧；不要卡死在 prompt 包。",
            "百炼额度/地域：Key 失效或欠费会导致全员出不了章，需值守账单。",
            "自训小模型：对比体量另排期；内测冻「百炼 + 本通道报告」，避免双轨口径混乱。",
        ],
    )
    add_heading(doc, "3.4 算法排期（对照多人内测）", 2)
    add_table(
        doc,
        ["窗口", "算法侧交付", "负责人建议"],
        [
            ["开测前 3 天", "A1–A3 全绿；固定一版出稿参数", "算法 + 后端联调"],
            ["开测前 1 天", "用 2 个真实邮箱各跑 1 日真图回归", "算法抽查文案/证据"],
            ["内测期", "失败 case 归类：空证据 / 模型超时 / 内容政策", "每日简报"],
            ["包后", "本地小模型 vs 百炼对比；评测集固化", "另轨"],
        ],
    )

    # ---- 4 ----
    add_heading(doc, "四、软件工作计划", 1)
    add_heading(doc, "4.1 已具备", 2)
    add_bullets(
        doc,
        [
            "App：入门、问卷、四 Tab、Evidence Lab、分享（拼贴/插画/Recap）、云同意开关。",
            "BFF：OTP、day-revisions、jobs、session Bearer、Dev Lite 与 Docker compose。",
            "本地轻处理：去重、768×1024、视频探针与选帧规则、远程 MobileCLIP 分析接口。",
            "文档：EXPO_GO_FULL_TEST_GUIDE（真机点法）、INTERNAL_BETA_FULL_GUIDE（A→D 内测总册）。",
        ],
    )
    add_heading(doc, "4.2 开多人内测前必须关的软件项", 2)
    add_table(
        doc,
        ["序号", "事项", "细节", "验收"],
        [
            [
                "S1",
                "公网 HTTPS BFF",
                "云主机或 Tunnel；反代到 8790；机器常开",
                "外网手机 Ping OK（非 192.168）",
            ],
            [
                "S2",
                "部署量产栈",
                "Postgres/Redis/MinIO/BFF/Worker + 同机 X02",
                "compose 健康；smoke-cloud-pack 过",
            ],
            [
                "S3",
                "换掉默认 API_TOKEN",
                "禁止群里长期共用 dev-token",
                "配置已轮换；旧 token 失效",
            ],
            [
                "S4",
                "OTP 全员可用",
                "SMTP 真发信；EXPOSE_OTP_IN_RESPONSE=false",
                "任意测试邮箱能收 6 位码",
            ],
            [
                "S5",
                "App/EAS 指公网",
                "EXPO_PUBLIC_BFF_BASE_URL=https://… 打进包或 Update",
                "内测员 Lab Apply 后 URL 正确",
            ],
            [
                "S6",
                "每用户日配额",
                "建议 5～10 次 day-revision/人/日；超限 429",
                "超限可测；不拖死机器",
            ],
            [
                "S7",
                "Worker 限并发",
                "并发 1～2，排队消化",
                "高峰排队而非雪崩",
            ],
            [
                "S8",
                "账号隔离抽检",
                "两邮箱各出一章，互不覆盖",
                "job/chapter 归属正确",
            ],
            [
                "S9",
                "分发包",
                "EAS preview → TestFlight（或固定 Expo 入口）",
                "5 人试跑后再扩到几十",
            ],
            [
                "S10",
                "值守与说明",
                "简短操作卡 + 失败回传模板",
                "群内有人盯 Ping/任务/账单",
            ],
        ],
    )
    add_heading(doc, "4.3 软件细节问题", 2)
    add_bullets(
        doc,
        [
            "BFF 地址踩坑：手机连不上 172.22（WSL）；本机联调必须 192.168.x；多人必须 https 公网。",
            "Dev Lite 与 Docker 抢 :8790：不要双开。",
            "start-dev-lite-background 默认 mock：UI 冒烟可以，质量验收不行。",
            "Android Expo Go 相册常被系统拦：内测说明引导用 iPhone 或 Lab A/A2。",
            "同意开关：Browse on this device only 必须拦住上传——回归必测。",
            "日记展示仍偏本地：云端保证 job 成功回写；若串号需补 job 绑 user_id。",
            "Recap 依赖 ffmpeg：公网机也要装，否则 Share day 失败。",
            "Lab / Skip 为内测捷径：商用包以后要藏；preview 可暂留 INTERNAL_BETA。",
            "防火墙与磁盘：MinIO/Postgres 爆盘会导致全员失败。",
        ],
    )
    add_heading(doc, "4.4 软件排期（建议 7 天立底座）", 2)
    add_table(
        doc,
        ["天", "软件交付", "完成标准"],
        [
            ["D0", "负责人自测 Expo Go 全量（阶段 A）", "手册勾选大半绿"],
            ["D1", "云主机 + compose + Key + mock off", "公网 /health；RequireGrounded 出章"],
            ["D2", "HTTPS + App 指向公网", "外网手机 Ping 通"],
            ["D3", "OTP 双账号隔离 + 换 API_TOKEN", "两邮箱各 1 chapter"],
            ["D4", "日配额补丁 + Worker 限并发", "超限明确错误"],
            ["D5", "TestFlight 或 Expo 入口 + 5 人试跑", "说明已发"],
            ["D6–7", "扩到几十人；盯账单与失败 job", "可重试；有值守"],
        ],
    )

    # ---- 5 ----
    add_heading(doc, "五、多人内测风险点（对照表）", 1)
    add_para(doc, "按「影响面 × 发生可能」管理；开测前红色项必须有对策。", bold=True)
    add_table(
        doc,
        ["等级", "风险", "表现", "对策"],
        [
            [
                "高",
                "仍用本机 Dev Lite 对多人",
                "外人连不上 / 你关机即全挂",
                "必须公网量产栈（S1/S2）",
            ],
            [
                "高",
                "mock 开着当质量过",
                "空日假章节，反馈失真",
                "ALLOW_MOCK_PREVIEW=false + 真图验收",
            ],
            [
                "高",
                "百炼 Key/额度爆",
                "全员任务失败",
                "配额 S6 + 账单值守 + 降并发",
            ],
            [
                "高",
                "无账号隔离 / 共用 token",
                "数据串、无法追责",
                "强制 OTP；轮换 API_TOKEN",
            ],
            [
                "中",
                "视频无 MobileCLIP",
                "选帧质量差或 degraded",
                "说明「照片主路径」或公网装权重",
            ],
            [
                "中",
                "Job 堆积 / 超时",
                "等很久或失败",
                "限并发、提示「可能数分钟」、可重试",
            ],
            [
                "中",
                "SMTP 进垃圾箱 / 发信失败",
                "登不进",
                "测前用多家邮箱试；备用管理员重发",
            ],
            [
                "中",
                "Android 相册失败",
                "部分用户无法导入",
                "优先 iOS；提供 Demo/A2 兜底说明",
            ],
            [
                "中",
                "公网机断电/磁盘满",
                "全员不可用",
                "云主机 + 磁盘告警 + 重启剧本",
            ],
            [
                "低",
                "UI 与 Figma 像素差",
                "观感吐槽",
                "锁定屏不大改；记日志不挡出章",
            ],
            [
                "低",
                "无 BLE",
                "Device 页未连接",
                "产品已定：诚实 Unavailable，不造假电量",
            ],
            [
                "低",
                "Expo Go 无 Vision",
                "Lab C 显示无原生 Vision",
                "视频靠 remote_bff 或降级；EAS 另轨",
            ],
        ],
    )

    # ---- 6 ----
    add_heading(doc, "六、联合里程碑与门禁", 1)
    add_table(
        doc,
        ["门禁", "含义", "算法", "软件", "未过能否对人"],
        [
            ["N1 链路绿", "相册→百炼→Diary 有章", "真 qwen", "公网或本机通", "仅内部自己"],
            ["N2 质量门", "无 mock；真 pack 有章", "RequireGrounded", "mock=false", "否"],
            ["N3 账号", "OTP 多用户", "—", "SMTP+session", "否（多人）"],
            ["N4 分享", "拼贴/Recap 可选", "—", "ffmpeg", "可降级开测"],
            ["N5 包", "TestFlight/入口", "—", "EAS preview", "否（普通人）"],
            ["N6 配额值守", "账单与堆积可控", "额度", "S6/S7/S10", "否"],
        ],
    )
    add_para(
        doc,
        "建议：N2+N3+N5+N6 全绿后再发「几十人」；N4 分享失败不挡 chapter 主目标。",
        bold=True,
    )

    # ---- 7 ----
    add_heading(doc, "七、内测执行与反馈机制", 1)
    add_heading(doc, "7.1 发给内测员的最短说明（要点）", 2)
    add_bullets(
        doc,
        [
            "安装 TestFlight（或指定 Expo 入口）；语言建议 English（对照手册时）。",
            "邮箱 OTP 登录 → Continue without a device → 问卷最后一屏选 Agree and continue。",
            "Privacy 打开 Allow cloud evidence upload。",
            "上传今天拍的 2～5 张真实照片；等待 Active tasks；去 Diary 看当天 chapter。",
            "出问题回传：Lab Ping 整行 + 任务失败原文 + 系统（iPhone/Android）。",
        ],
    )
    add_heading(doc, "7.2 内部日报字段", 2)
    add_bullets(
        doc,
        [
            "今日成功 chapter 数 / 失败 job 数 / 主要错误码。",
            "百炼调用量与费用粗估。",
            "Top 3 用户吐槽（导入失败、太慢、空日、登录收不到码）。",
            "是否触发配额、是否需扩容或降并发。",
        ],
    )
    add_heading(doc, "7.3 回滚与降级开关", 2)
    add_bullets(
        doc,
        [
            "算法挂：临时允许只收素材、提示「生成延迟」（勿偷偷开 mock 假装有记忆）。",
            "视频语义挂：强制照片路径或时间窗选帧。",
            "公网挂：公告暂停上传；保留本地浏览。",
            "额度将尽：下调日配额或暂停新用户。",
        ],
    )

    # ---- 8 ----
    add_heading(doc, "八、责任分工建议", 1)
    add_table(
        doc,
        ["角色", "主责"],
        [
            ["产品/项目", "内测范围、说明文案、是否放宽视频/分享"],
            ["软件客户端", "App 指公网、TestFlight、同意拦截、失败展示"],
            ["软件后端", "公网部署、OTP、配额、Worker、健康检查"],
            ["算法", "关 mock、真 pack 质量、失败归类、参数冻结"],
            ["值守", "账单、磁盘、失败 job、群内答疑"],
        ],
    )

    # ---- 9 ----
    add_heading(doc, "九、一页纸结论（汇报用）", 1)
    add_para(
        doc,
        "算法侧：合同与通道已有，开测前把「真千问 + 关 mock + 至少一份真实 grounded pack」做实；视频 MobileCLIP 可降级，照片是主验收。",
    )
    add_para(
        doc,
        "软件侧：App 与本机链路已能演示；多人成立的前提是「公网量产 BFF + OTP 隔离 + 粗配额 + TestFlight/固定入口」，缺一不可。",
    )
    add_para(
        doc,
        "最大风险：用家里 Dev Lite 假装多人服务、用 mock 章节假装质量过关、无配额打爆百炼。按阶段 C→D 关门禁后再放量。",
        bold=True,
    )
    add_para(
        doc,
        "配套手册：docs/INTERNAL_BETA_FULL_GUIDE.md（总册）、docs/EXPO_GO_FULL_TEST_GUIDE.md（真机点法）、docs/CLOUD_BULK_CHANNEL.md（PC 质量通道）、docs/WINDOWS_MOBILECLIP_BFF.md（向量权重）。",
        size=10,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
